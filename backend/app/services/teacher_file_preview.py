"""Read DOCX content without executing document code or exposing external links."""
from zipfile import ZipFile, BadZipFile
from xml.etree import ElementTree as ET
from pathlib import Path
import re

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_preview(path):
    try:
        with ZipFile(path) as archive:
            info = archive.getinfo("word/document.xml")
            if info.file_size > 20 * 1024 * 1024:
                raise ValueError("Documento demasiado grande para la vista previa.")
            xml = archive.read(info)
        if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
            raise ValueError("Documento no válido.")
        root = ET.fromstring(xml)
        body = root.find(W + "body")
        if body is None:
            raise ValueError("Documento sin contenido.")
    except (BadZipFile, KeyError, ET.ParseError, OSError) as exc:
        raise ValueError("No se puede leer este documento Word.") from exc

    def paragraph(node):
        style = node.find("./" + W + "pPr/" + W + "pStyle")
        style_name = style.get(W + "val", "") if style is not None else ""
        heading = next((n for n in range(1, 7) if style_name.lower() in (f"heading{n}", f"titulo{n}", f"título{n}")), 0)
        runs = []
        for run in node.iter(W + "r"):
            text = "".join(item.text or "" if item.tag == W + "t" else "\n" if item.tag == W + "br" else "\t" if item.tag == W + "tab" else "" for item in run)
            if not text:
                continue
            props = run.find(W + "rPr")
            def active(name):
                prop = props.find(W + name) if props is not None else None
                return prop is not None and prop.get(W + "val", "true") not in ("0", "false", "off", "none")
            runs.append({"text": text, "bold": active("b"), "italic": active("i"), "underline": active("u")})
        return {"type": "paragraph", "heading": heading, "list": node.find("./" + W + "pPr/" + W + "numPr") is not None, "runs": runs}

    blocks = []
    for node in body:
        if node.tag == W + "p":
            blocks.append(paragraph(node))
        elif node.tag == W + "tbl":
            blocks.append({"type": "table", "rows": [[ [paragraph(p) for p in cell.findall(W + "p")] for cell in row.findall(W + "tc")] for row in node.findall(W + "tr")]})
    return {"format": "docx", "blocks": blocks}


def _plain_block(text, heading=0):
    return {"type": "paragraph", "heading": heading, "list": False, "runs": [{"text": str(text)}]}


def _safe_zip_xml(archive, name):
    info = archive.getinfo(name)
    if info.file_size > 20 * 1024 * 1024:
        raise ValueError("Documento demasiado grande para la vista previa.")
    xml = archive.read(info)
    if b"<!DOCTYPE" in xml or b"<!ENTITY" in xml:
        raise ValueError("Documento no válido.")
    return ET.fromstring(xml)


def _check_zip_size(path):
    with ZipFile(path) as archive:
        files = archive.infolist()
        if len(files) > 2000 or sum(info.file_size for info in files) > 80 * 1024 * 1024:
            raise ValueError("Documento demasiado grande para la vista previa.")


def document_preview(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        return docx_preview(path)
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        try:
            _check_zip_size(path)
            workbook = load_workbook(path, read_only=True, data_only=True)
            blocks = []
            for sheet in workbook.worksheets[:20]:
                blocks.append(_plain_block(sheet.title, 2))
                rows = []
                if sheet.max_row and sheet.max_column:
                    for row in sheet.iter_rows(max_row=min(sheet.max_row, 200), max_col=min(sheet.max_column, 30), values_only=True):
                        rows.append([[_plain_block(value if value is not None else "")] for value in row])
                if rows:
                    blocks.append({"type": "table", "rows": rows})
            workbook.close()
            return {"format": "xlsx", "blocks": blocks}
        except (OSError, ValueError, BadZipFile, KeyError) as exc:
            raise ValueError("No se puede leer esta hoja de cálculo.") from exc
    if suffix in {".pptx", ".odt"}:
        try:
            _check_zip_size(path)
            with ZipFile(path) as archive:
                if suffix == ".odt":
                    roots = [_safe_zip_xml(archive, "content.xml")]
                else:
                    names = sorted((name for name in archive.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)), key=lambda name: int(re.search(r"\d+", name).group()))[:100]
                    roots = [_safe_zip_xml(archive, name) for name in names]
            blocks = []
            for index, root in enumerate(roots):
                if suffix == ".pptx":
                    blocks.append(_plain_block(f"Diapositiva {index + 1}", 2))
                    texts = [node.text or "" for node in root.iter() if node.tag.endswith("}t") and node.text]
                    blocks.extend(_plain_block(text) for text in texts[:300])
                else:
                    for node in root.iter():
                        if node.tag.endswith(("}p", "}h")):
                            text = "".join(node.itertext()).strip()
                            if text:
                                blocks.append(_plain_block(text, 2 if node.tag.endswith("}h") else 0))
                            if len(blocks) >= 500:
                                break
            return {"format": suffix[1:], "blocks": blocks}
        except (OSError, ValueError, BadZipFile, KeyError, ET.ParseError) as exc:
            raise ValueError("No se puede leer este documento.") from exc
    if suffix == ".rtf":
        raw = Path(path).read_bytes()
        if not raw.lstrip().startswith(b"{\\rtf"):
            raise ValueError("El documento RTF no es válido.")
        text = raw.decode("latin-1", errors="replace")
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
        text = re.sub(r"\\[^a-zA-Z]", " ", text).replace("{", "").replace("}", "")
        return {"format": "rtf", "blocks": [_plain_block(text[:100_000].strip())]}
    raise ValueError("Este formato no dispone de vista previa.")
