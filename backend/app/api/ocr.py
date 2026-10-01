from io import BytesIO
from pathlib import Path
from uuid import UUID, uuid4
import shutil
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import FileResponse, RedirectResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.orm import Session
from app.auth.dependencies import current_user
from app.database.session import get_db
from app.models import User, UsageRecord
from app.models.ocr import OcrJob
from app.services.licenses import check_access
from app.services.ocr_storage import job_dir, model_config

router = APIRouter(tags=["ocr"])
MAX_BYTES = 15 * 1024 * 1024
FIELDS = ("proveedor", "nif_cif", "empresa_receptora", "numero_factura", "fecha_factura", "base_imponible", "iva_porcentaje", "iva_importe", "irpf_porcentaje", "irpf_importe", "total_pagar", "moneda", "concepto_mejorado")


def authorized(request: Request, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if request.method != "GET" and request.headers.get("sec-fetch-site") == "cross-site":
        raise HTTPException(403, "Solicitud no permitida")
    decision = check_access(db, user, "OCR_FACTURAS", require_credits=False)
    if not decision.ok:
        raise HTTPException(403, decision.message)
    return user


def owned(db, user, job_id):
    job = db.scalar(select(OcrJob).where(OcrJob.id == job_id, OcrJob.user_id == user.id, OcrJob.organization_id == user.organization_id))
    if not job:
        raise HTTPException(404, "Factura no encontrada")
    return job


def summary(job):
    return {"id": str(job.id), "filename": job.filename, "status": job.status, "error": job.error, "result": job.result, "created_at": job.created_at.isoformat()}


@router.get("/facturas")
def page(request: Request, db: Session = Depends(get_db)):
    from app.api.app_routes import page_user_or_redirect, static_html, PROJECT_ROOT
    user = page_user_or_redirect(request, db, "/facturas")
    if isinstance(user, RedirectResponse):
        return user
    return static_html("ocr-facturas-app.html", PROJECT_ROOT / "apps" / "ocr" / "index.html")


@router.get("/api/ocr/status")
def status(user: User = Depends(authorized), db: Session = Depends(get_db)):
    try:
        cfg = model_config(db)
        return {"configured": True, "provider": cfg["api_tipo"], "model": cfg["modelo_externo"]}
    except ValueError as exc:
        return {"configured": False, "message": str(exc)}


@router.get("/api/ocr/jobs")
def jobs(user: User = Depends(authorized), db: Session = Depends(get_db)):
    rows = db.scalars(select(OcrJob).where(OcrJob.user_id == user.id, OcrJob.organization_id == user.organization_id).order_by(OcrJob.created_at.desc()).limit(100)).all()
    return {"jobs": [summary(row) for row in rows]}


def validate_document(data, suffix):
    try:
        if suffix == ".pdf":
            import fitz
            with fitz.open(stream=data, filetype="pdf") as doc:
                if doc.needs_pass or not 1 <= len(doc) <= 10:
                    raise ValueError()
        else:
            from PIL import Image
            with Image.open(BytesIO(data)) as img:
                if img.format not in {"PNG", "JPEG", "WEBP"} or img.width * img.height > 25_000_000:
                    raise ValueError()
                img.verify()
    except Exception:
        raise HTTPException(422, "Documento no válido. Máximo 10 páginas PDF sin contraseña o imagen de 25 megapíxeles.")


@router.post("/api/ocr/jobs", status_code=202)
async def upload(request: Request, filename: str, user: User = Depends(authorized), db: Session = Depends(get_db)):
    if request.headers.get("content-type", "").split(";")[0] != "application/octet-stream":
        raise HTTPException(415, "Envía el documento como application/octet-stream")
    name = filename.replace("\\", "/").split("/")[-1][:180]
    suffix = Path(name).suffix.lower()
    if suffix not in {".pdf", ".png", ".jpg", ".jpeg", ".webp"}:
        raise HTTPException(415, "Formato admitido: PDF, PNG, JPG o WEBP")
    try:
        config = model_config(db)
    except ValueError as exc:
        raise HTTPException(503, str(exc))
    data = bytearray()
    async for chunk in request.stream():
        data.extend(chunk)
        if len(data) > MAX_BYTES:
            raise HTTPException(413, "Máximo 15 MB por documento")
    if not data:
        raise HTTPException(422, "El documento está vacío")
    validate_document(bytes(data), suffix)
    # Serialize credit reservations for this user, including parallel uploads.
    db.scalar(select(User).where(User.id == user.id).with_for_update())
    decision = check_access(db, user, "OCR_FACTURAS", require_credits=True)
    if not decision.ok:
        raise HTTPException(402, decision.message)
    active = db.scalar(select(func.count()).select_from(OcrJob).where(OcrJob.user_id == user.id, OcrJob.status.in_(["queued", "processing"])))
    if active >= 3:
        raise HTTPException(429, "Ya tienes tres documentos en curso. Espera a que terminen.")
    usage = UsageRecord(id=uuid4(), organization_id=user.organization_id, user_id=user.id, product_code="OCR_FACTURAS", model=config["modelo_externo"])
    job = OcrJob(id=uuid4(), organization_id=user.organization_id, user_id=user.id, usage_id=usage.id, filename=name, suffix=suffix, provider=config["api_tipo"], model=config["modelo_externo"], status="queued", error="", result={})
    folder = job_dir(job)
    try:
        folder.mkdir(parents=True, exist_ok=False)
        (folder / ("source" + suffix)).write_bytes(data)
        db.add(usage)
        db.flush()
        db.add(job)
        db.commit()
    except Exception:
        db.rollback()
        shutil.rmtree(folder, ignore_errors=True)
        raise
    return summary(job)


@router.get("/api/ocr/jobs/{job_id}")
def detail(job_id: UUID, user: User = Depends(authorized), db: Session = Depends(get_db)):
    return summary(owned(db, user, job_id))


class Review(BaseModel):
    fields: dict[str, str] = Field(max_length=20)


@router.post("/api/ocr/jobs/{job_id}/review")
def review(job_id: UUID, body: Review, user: User = Depends(authorized), db: Session = Depends(get_db)):
    job = owned(db, user, job_id)
    if job.status != "done":
        raise HTTPException(409, "Espera a que termine el procesamiento")
    if any(key not in FIELDS or len(value) > 2000 for key, value in body.fields.items()):
        raise HTTPException(422, "Campos no válidos")
    job.result = {**job.result, "fields": {**job.result.get("fields", {}), **body.fields}, "reviewed": True}
    db.commit()
    return summary(job)


@router.get("/api/ocr/jobs/{job_id}/original")
def original(job_id: UUID, user: User = Depends(authorized), db: Session = Depends(get_db)):
    job = owned(db, user, job_id)
    return FileResponse(job_dir(job) / ("source" + job.suffix), filename=job.filename, media_type="application/octet-stream", headers={"X-Content-Type-Options": "nosniff", "Cache-Control": "no-store"})


@router.get("/api/ocr/jobs/{job_id}/export")
def export(job_id: UUID, user: User = Depends(authorized), db: Session = Depends(get_db)):
    from openpyxl import Workbook
    job = owned(db, user, job_id)
    if job.status != "done":
        raise HTTPException(409, "Resultado no disponible")
    wb = Workbook()
    ws = wb.active
    ws.title = "Factura"
    ws.append(["Campo", "Valor"])
    for key in FIELDS:
        value = str(job.result.get("fields", {}).get(key, ""))
        # Explicit text type prevents invoice content becoming spreadsheet formulas.
        ws.append([key, value])
        ws.cell(ws.max_row, 2).data_type = "s"
    ws.append(["Revisada", "Sí" if job.result.get("reviewed") else "No"])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 55
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return StreamingResponse(stream, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="factura.xlsx"', "Cache-Control": "no-store"})
